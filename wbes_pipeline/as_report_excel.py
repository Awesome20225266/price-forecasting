"""
Excel workbook for as_report.py (kept separate so the report logic stays readable).

Layout rule used on every table sheet: title in row 1, a one-line note in row 2, column headers in
row 4, data from row 5. Data sheets ("Daily by Plant", "Hourly Data", "Matrix Up/Down", "Data Check")
hold values read from the source files; every summary sheet is formulas over those data sheets, so
the totals stay traceable and the "Range Analyser" recalculates when its yellow input cells change.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
UP_HEX, DOWN_HEX = "2A78D6", "EB6834"           # same two colours as the HTML dashboard
F_BASE = Font(name=FONT, size=10)
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
F_SECTION = Font(name=FONT, size=11, bold=True, color="1F3864")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_INPUT = Font(name=FONT, size=10, bold=True, color="0000FF")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
HEAD_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")

MWH = '#,##0.0;-#,##0.0;"-"'
MWH0 = '#,##0;-#,##0;"-"'
MU = '#,##0.00" MU";-#,##0.00" MU";"-"'
PCT = '0.0%;-0.0%;"-"'
INT = '#,##0;-#,##0;"-"'
DATE = "dd-mmm-yyyy"

FIRST = 5                                        # first data row on every table sheet


# ----------------------------------------------------------------------------- small helpers
def put(ws, r: int, c: int, value, fmt: str | None = None, font: Font | None = None, fill=None, align=None):
    cell = ws.cell(row=r, column=c, value=value)
    if fmt:
        cell.number_format = fmt
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    return cell


def titled(ws, title: str, note: str) -> None:
    put(ws, 1, 1, title, font=F_TITLE)
    put(ws, 2, 1, note, font=F_NOTE)


def header(ws, r: int, names: list[str], widths: list[float] | None = None, c0: int = 1) -> None:
    for i, name in enumerate(names):
        put(ws, r, c0 + i, name, font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
        if widths:
            ws.column_dimensions[col(c0 + i)].width = widths[i]
    ws.row_dimensions[r].height = 30


def _num(v):
    """Float for a cell, None (blank) for NaN / missing."""
    return None if v is None or v != v else float(v)


def ref(sheet: str, c: str, r1: int, r2: int) -> str:
    return f"'{sheet}'!${c}${r1}:${c}${r2}"


def total_row(ws, r: int, c_from: int, c_to: int) -> None:
    for c in range(c_from, c_to + 1):
        ws.cell(row=r, column=c).fill = FILL_TOTAL
        ws.cell(row=r, column=c).font = F_BOLD


def bar_chart(title: str, y_title: str, cats: Reference, series: list[tuple[Reference, str, str]],
              width: float = 22, height: float = 9) -> BarChart:
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth, ch.overlap = "col", "clustered", 40, -8
    ch.title = f"{title} ({y_title})"
    for data, name, colour in series:
        ch.add_data(data, titles_from_data=False)
        s = ch.series[-1]
        s.tx = SeriesLabel(v=name)
        s.graphicalProperties.solidFill = colour
        s.graphicalProperties.line.noFill = True
    ch.set_categories(cats)
    _axes(ch)
    ch.width, ch.height = width, height
    return ch


def line_chart(title: str, y_title: str, cats: Reference, series: list[tuple[Reference, str, str]],
               width: float = 22, height: float = 9) -> LineChart:
    ch = LineChart()
    ch.title = f"{title} ({y_title})"
    for data, name, colour in series:
        ch.add_data(data, titles_from_data=False)
        s = ch.series[-1]
        s.tx = SeriesLabel(v=name)
        s.graphicalProperties.line.solidFill = colour
        s.graphicalProperties.line.width = 19050      # 1.5 pt
        s.smooth = False
        s.marker.symbol = "none"
    ch.set_categories(cats)
    ch.x_axis.number_format = "dd-mmm"
    _axes(ch)
    ch.width, ch.height = width, height
    return ch


def _axes(ch) -> None:
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = "#,##0"
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E0E0E0"))
    ch.legend.position = "t"
    ch.legend.overlay = False
    ch.title.overlay = False


# ----------------------------------------------------------------------------- workbook
def write_excel(t: dict, path: Path, bands: list, regions: list[str], tech_order: list[str]) -> Path:
    day: pd.DataFrame = t["day"]
    plants: pd.DataFrame = t["plants"]
    dates = [pd.Timestamp(d).to_pydatetime() for d in t["dates"]]
    months = sorted(day["month"].unique())
    types = list(dict.fromkeys(plants["type"]))
    techs = [x for x in tech_order if x in set(plants["technology"])]
    band_names = [b for b, _ in bands]
    nb = len(bands)

    wb = Workbook()
    wb._fonts = IndexedList([F_BASE])                 # Arial 10 as the workbook default font
    wb._named_styles["Normal"].font = F_BASE
    sheets = ["Read Me", "Summary", "Plant Summary", "Plants", "Daily by Plant", "Schedule by Type", "Daily Totals", "Matrix Up",
              "Matrix Down", "Time Slots", "Plant Time Bands", "Range Analyser", "Data Check", "Hourly Data",
              "Lists"]
    ws = {name: (wb.active if i == 0 else wb.create_sheet()) for i, name in enumerate(sheets)}
    for name, w in ws.items():
        w.title = name

    n_day, n_pl, n_dt = len(day), len(plants), len(dates)
    type_col = col(16 + 2 * nb)                       # 'Daily by Plant': WBES plant type, after the band columns
    DL = FIRST + n_day - 1                            # last data row, Daily by Plant / Hourly Data
    PL = FIRST + n_pl - 1                             # last data row, Plant Summary / bands / matrices
    TL = FIRST + n_dt - 1                             # last data row, Daily Totals
    P_TOT, T_TOT = PL + 2, TL + 2                     # total rows (one blank row after the data)

    DP = lambda c: ref("Daily by Plant", c, FIRST, DL)
    HD = lambda c: ref("Hourly Data", c, FIRST, DL)
    PS = lambda c: ref("Plant Summary", c, FIRST, PL)
    PB = lambda c: ref("Plant Time Bands", c, FIRST, PL)
    sched, stypes = t["sched"], t["stypes"]
    SL = FIRST + len(sched) - 1                       # last data row, Schedule by Type
    SB = lambda c: ref("Schedule by Type", c, FIRST, SL)
    st_col = {x: col(6 + i) for i, x in enumerate(stypes)}  # Schedule by Type: one column per schedule type

    # ---------------------------------------------------------------- Daily by Plant (data)
    s = ws["Daily by Plant"]
    titled(s, "AS schedule per plant per day",
           "One row per plant per day with any AS. Up = extra injection, Down = backing down (MWh = MW x 0.25 h). "
           "Time slots = 15-min blocks with AS that day. Values read from wbes_ns_data; Net and Gross are formulas.")
    band_up = [f"up|{b}" for b in band_names]
    band_dn = [f"down|{b}" for b in band_names]
    header(s, 4, ["Date", "Month", "Region", "Plant", "Technology", "AS Up (MWh)", "AS Down (MWh)",
                  "Net Up-Down (MWh)", "Gross Up+Down (MWh)", "Up blocks", "Down blocks", "Peak Up (MW)",
                  "Peak Down (MW)", "Up time slots", "Down time slots"]
           + [f"Up MWh {b}" for b in band_names] + [f"Down MWh {b}" for b in band_names]
           + ["WBES plant type", "Total schedule that day (MWh, + = into grid)", "AS net as % of Total"],
           [12, 9, 8, 26, 22, 12, 12, 12, 12, 8, 8, 11, 11, 45, 45] + [13] * (2 * nb) + [22, 14, 11])
    r = FIRST
    for row in day[["date", "month", "region", "plant", "technology", "up_mwh", "down_mwh", "up_blocks",
                    "down_blocks", "peak_up_mw", "peak_down_mw", "up_slots", "down_slots"]
                   + band_up + band_dn + ["type", "total_mwh", "as_share"]].itertuples(index=False):
        put(s, r, 1, row[0].to_pydatetime(), DATE)
        for c, v in ((2, row[1]), (3, row[2]), (4, row[3]), (5, row[4])):
            s.cell(row=r, column=c, value=v)
        put(s, r, 6, float(row[5]), MWH)
        put(s, r, 7, float(row[6]), MWH)
        put(s, r, 8, f"=F{r}-G{r}", MWH)
        put(s, r, 9, f"=F{r}+G{r}", MWH)
        put(s, r, 10, int(row[7]), INT)
        put(s, r, 11, int(row[8]), INT)
        put(s, r, 12, float(row[9]), MWH)
        put(s, r, 13, float(row[10]), MWH)
        s.cell(row=r, column=14, value=row[11] or None)
        s.cell(row=r, column=15, value=row[12] or None)
        for i, v in enumerate(row[13:13 + 2 * nb]):
            put(s, r, 16 + i, float(v), MWH)
        s.cell(row=r, column=16 + 2 * nb, value=row[-3])
        put(s, r, 17 + 2 * nb, _num(row[-2]), MWH)
        put(s, r, 18 + 2 * nb, _num(row[-1]), PCT)
        r += 1
    s.freeze_panes = "E5"
    s.auto_filter.ref = f"A4:{col(18 + 2 * nb)}{DL}"

    # ---------------------------------------------------------------- Schedule by Type (data)
    s = ws["Schedule by Type"]
    titled(s, "Whole schedule of every plant that had AS - MWh per schedule type, every day",
           "+ = energy into the grid (injection), - = drawn from the grid (e.g. BESS charging). AS here is net "
           "(Up - Down). The types add up to Total in every row (checked). Includes days without AS.")
    as_c, tot_c = st_col.get("AS"), st_col.get("Total")
    header(s, 4, ["Date", "Month", "Region", "Plant", "Technology"] + [f"{x} (MWh)" for x in stypes]
           + ["AS net as % of Total"], [12, 9, 8, 26, 22] + [12] * len(stypes) + [11])
    for r, row in zip(range(FIRST, SL + 1), sched[["date", "month", "region", "plant", "technology"]
                                                   + stypes].itertuples(index=False)):
        put(s, r, 1, row[0].to_pydatetime(), DATE)
        for c in range(1, 5):
            s.cell(row=r, column=c + 1, value=row[c])
        for i, v in enumerate(row[5:]):
            if abs(v) > 1e-9:
                put(s, r, 6 + i, float(v), MWH)
        if as_c and tot_c:
            put(s, r, 6 + len(stypes), f'=IF(ABS(N({tot_c}{r}))<0.000001,"",N({as_c}{r})/{tot_c}{r})', PCT)
    s.freeze_panes = "F5"
    s.auto_filter.ref = f"A4:{col(6 + len(stypes))}{SL}"

    # ---------------------------------------------------------------- Hourly Data (data)
    s = ws["Hourly Data"]
    titled(s, "AS energy per plant per day per hour (MWh) - data for the time-slot sheets",
           "Same rows as 'Daily by Plant'. Hour 00 = 00:00-01:00. Blank = no AS in that hour. Up and Down are "
           "positive magnitudes.")
    up_cols = [f"up_h{h:02d}" for h in range(24)]
    dn_cols = [f"down_h{h:02d}" for h in range(24)]
    header(s, 4, ["Date", "Month", "Region", "Plant", "Technology"]
           + [f"Up {h:02d}-{h + 1:02d}" for h in range(24)] + [f"Down {h:02d}-{h + 1:02d}" for h in range(24)],
           [12, 9, 8, 26, 22] + [9] * 48)
    keys = day[["date", "month", "region", "plant", "technology"]].itertuples(index=False)
    for r, k, vals in zip(range(FIRST, DL + 1), keys, day[up_cols + dn_cols].to_numpy()):
        put(s, r, 1, k[0].to_pydatetime(), DATE)
        for c in range(1, 5):
            s.cell(row=r, column=c + 1, value=k[c])
        for i, v in enumerate(vals):
            if v > 0:
                put(s, r, 6 + i, float(v), MWH)
    s.freeze_panes = "F5"
    s.auto_filter.ref = f"A4:{col(53)}{DL}"
    hour_col = lambda h, up=True: col(6 + h if up else 30 + h)

    # ---------------------------------------------------------------- Plant Summary
    s = ws["Plant Summary"]
    titled(s, "Cumulative AS per plant over the whole date range",
           f"Formulas over 'Daily by Plant'. Sorted by Gross (Up+Down) at the time of generation; "
           f"use the filter buttons to re-sort. Main time band = band with the most MWh ('Plant Time Bands').")
    header(s, 4, ["Rank (by Gross)", "Region", "Plant", "Technology", "Days with AS", "Days with Up",
                  "Days with Down", "AS Up (MWh)", "AS Down (MWh)", "Net Up-Down (MWh)", "Gross Up+Down (MWh)",
                  "Share of all-plant Gross", "Avg Gross per AS day (MWh)", "Peak Up (MW)", "Peak Down (MW)",
                  "First AS date", "Last AS date", "Main Up time band", "Main Down time band",
                  "Up sort key", "Down sort key", "WBES plant type", "Max scheduled MW (WBES)",
                  "Nameplate MW (web)", "Nameplate source"],
           [8, 8, 26, 22, 9, 9, 9, 13, 13, 13, 13, 10, 12, 10, 10, 12, 12, 20, 20, 6, 6, 20, 11, 11, 40])
    band_rng = ref("Time Slots", "A", 35, 34 + len(bands))
    for i, p in enumerate(plants.itertuples(index=False)):
        r = FIRST + i
        put(s, r, 1, f"=RANK(K{r},$K${FIRST}:$K${PL})", INT)
        s.cell(row=r, column=2, value=p.region)
        s.cell(row=r, column=3, value=p.plant)
        s.cell(row=r, column=4, value=p.technology)
        s.cell(row=r, column=22, value=p.type)
        put(s, r, 23, _num(p.max_sched_mw), MWH)
        put(s, r, 24, _num(p.nameplate_mw), MWH)
        s.cell(row=r, column=25, value=p.source_url or None)
        put(s, r, 5, f"=COUNTIFS({DP('D')},$C{r})", INT)
        put(s, r, 6, f"=COUNTIFS({DP('D')},$C{r},{DP('F')},\">0\")", INT)
        put(s, r, 7, f"=COUNTIFS({DP('D')},$C{r},{DP('G')},\">0\")", INT)
        put(s, r, 8, f"=SUMIFS({DP('F')},{DP('D')},$C{r})", MWH)
        put(s, r, 9, f"=SUMIFS({DP('G')},{DP('D')},$C{r})", MWH)
        put(s, r, 10, f"=H{r}-I{r}", MWH)
        put(s, r, 11, f"=H{r}+I{r}", MWH)
        put(s, r, 12, f"=IF($K${P_TOT}=0,0,K{r}/$K${P_TOT})", PCT)
        put(s, r, 13, f"=IF(E{r}=0,0,K{r}/E{r})", MWH)
        put(s, r, 14, f"=_xlfn.MAXIFS({DP('L')},{DP('D')},$C{r})", MWH)
        put(s, r, 15, f"=_xlfn.MAXIFS({DP('M')},{DP('D')},$C{r})", MWH)
        put(s, r, 16, f"=_xlfn.MINIFS({DP('A')},{DP('D')},$C{r})", DATE)
        put(s, r, 17, f"=_xlfn.MAXIFS({DP('A')},{DP('D')},$C{r})", DATE)
        for c, (c1, c2) in ((18, (col(4), col(3 + nb))), (19, (col(5 + 2 * nb), col(4 + 3 * nb)))):
            row_vals = (f"INDEX('Plant Time Bands'!${c1}${FIRST}:${c2}${PL},"
                        f"MATCH($C{r},{PB('B')},0),0)")
            s.cell(row=r, column=c, value=f"=IFERROR(IF(MAX({row_vals})=0,\"-\","
                                          f"INDEX({band_rng},MATCH(MAX({row_vals}),{row_vals},0))),\"-\")")
        put(s, r, 20, f"=H{r}+ROW()/1E9", "0.0")       # unique keys for the Top-10 lists on Summary
        put(s, r, 21, f"=I{r}+ROW()/1E9", "0.0")
    r = P_TOT
    put(s, r, 3, "All plants")
    for c, f in ((5, "SUM"), (6, "SUM"), (7, "SUM"), (8, "SUM"), (9, "SUM"), (10, "SUM"), (11, "SUM"),
                 (12, "SUM"), (14, "MAX"), (15, "MAX")):
        L = col(c)
        put(s, r, c, f"={f}({L}{FIRST}:{L}{PL})", PCT if c == 12 else (INT if c < 8 else MWH))
    total_row(s, r, 1, 19)
    s.freeze_panes = "D5"
    c0 = 26
    put(s, 3, c0, "Whole schedule over the period, MWh by type (+ = into grid, - = drawn from grid; AS is net)",
        font=F_SECTION)
    header(s, 4, [f"{x} (MWh)" for x in stypes] + ["AS net as % of Total"], [12] * len(stypes) + [11], c0=c0)
    for i in range(n_pl):
        r = FIRST + i
        for k, x in enumerate(stypes):
            put(s, r, c0 + k, f"=SUMIFS({SB(st_col[x])},{SB('D')},$C{r})", MWH)
        if "AS" in stypes and "Total" in stypes:
            a_, t_ = col(c0 + stypes.index("AS")), col(c0 + stypes.index("Total"))
            put(s, r, c0 + len(stypes), f'=IF(ABS({t_}{r})<0.000001,"",{a_}{r}/{t_}{r})', PCT)
    for k in range(len(stypes)):
        L = col(c0 + k)
        put(s, P_TOT, c0 + k, f"=SUM({L}{FIRST}:{L}{PL})", MWH)
    total_row(s, P_TOT, c0, c0 + len(stypes))
    s.auto_filter.ref = f"A4:{col(c0 + len(stypes))}{PL}"
    s.column_dimensions["T"].hidden = s.column_dimensions["U"].hidden = True

    # ---------------------------------------------------------------- Plant Time Bands
    s = ws["Plant Time Bands"]
    titled(s, "When does each plant get its AS? MWh and share by time band, whole date range",
           "Formulas over the band columns of 'Daily by Plant'. Bands: "
           + "; ".join(band_names) + " (hour the 15-min block starts in).")
    groups = [("AS Up - MWh by time band", 4), ("AS Up - share of the plant's Up", 5 + nb),
              ("AS Down - MWh by time band", 5 + 2 * nb), ("AS Down - share of the plant's Down", 6 + 3 * nb)]
    for text, c in groups:
        put(s, 3, c, text, font=F_SECTION)
    header(s, 4, ["Region", "Plant", "Technology"] + band_names + ["Up total (MWh)"] + band_names
           + band_names + ["Down total (MWh)"] + band_names,
           [8, 26, 22] + [12] * nb + [13] + [10] * nb + [12] * nb + [13] + [10] * nb)
    for i, p in enumerate(plants.itertuples(index=False)):
        r = FIRST + i
        s.cell(row=r, column=1, value=p.region)
        s.cell(row=r, column=2, value=p.plant)
        s.cell(row=r, column=3, value=p.technology)
        for k in range(nb):
            put(s, r, 4 + k, f"=SUMIFS({DP(col(16 + k))},{DP('D')},$B{r})", MWH)
            put(s, r, 5 + 2 * nb + k, f"=SUMIFS({DP(col(16 + nb + k))},{DP('D')},$B{r})", MWH)
        up_tot, dn_tot = col(4 + nb), col(5 + 3 * nb)
        put(s, r, 4 + nb, f"=SUM({col(4)}{r}:{col(3 + nb)}{r})", MWH)
        put(s, r, 5 + 3 * nb, f"=SUM({col(5 + 2 * nb)}{r}:{col(4 + 3 * nb)}{r})", MWH)
        for k in range(nb):
            put(s, r, 5 + nb + k, f"=IF(${up_tot}{r}=0,0,{col(4 + k)}{r}/${up_tot}{r})", PCT)
            put(s, r, 6 + 3 * nb + k, f"=IF(${dn_tot}{r}=0,0,{col(5 + 2 * nb + k)}{r}/${dn_tot}{r})", PCT)
    r = P_TOT
    put(s, r, 2, "All plants")
    for c in list(range(4, 5 + nb)) + list(range(5 + 2 * nb, 6 + 3 * nb)):
        put(s, r, c, f"=SUM({col(c)}{FIRST}:{col(c)}{PL})", MWH)
    for k in range(nb):
        put(s, r, 5 + nb + k, f"=IF(${col(4 + nb)}{r}=0,0,{col(4 + k)}{r}/${col(4 + nb)}{r})", PCT)
        put(s, r, 6 + 3 * nb + k, f"=IF(${col(5 + 3 * nb)}{r}=0,0,{col(5 + 2 * nb + k)}{r}/${col(5 + 3 * nb)}{r})", PCT)
    total_row(s, r, 1, 5 + 4 * nb)
    s.freeze_panes = "D5"
    s.auto_filter.ref = f"A4:{col(5 + 4 * nb)}{PL}"

    # ---------------------------------------------------------------- Daily Totals
    s = ws["Daily Totals"]
    titled(s, "All plants together, per day and per month",
           "Formulas over 'Daily by Plant'. Region columns split the same totals by region.")
    header(s, 4, ["Date", "Month", "Plants with AS", "AS Up (MWh)", "AS Down (MWh)", "Net Up-Down (MWh)",
                  "Gross Up+Down (MWh)"] + [f"Up {g} (MWh)" for g in regions] + [f"Down {g} (MWh)" for g in regions],
           [12, 9, 9, 13, 13, 13, 13] + [11] * (2 * len(regions)))
    for i, d in enumerate(dates):
        r = FIRST + i
        put(s, r, 1, d, DATE)
        s.cell(row=r, column=2, value=f"{d:%Y-%m}")
        put(s, r, 3, f"=COUNTIFS({DP('A')},$A{r})", INT)
        put(s, r, 4, f"=SUMIFS({DP('F')},{DP('A')},$A{r})", MWH)
        put(s, r, 5, f"=SUMIFS({DP('G')},{DP('A')},$A{r})", MWH)
        put(s, r, 6, f"=D{r}-E{r}", MWH)
        put(s, r, 7, f"=D{r}+E{r}", MWH)
        for k, g in enumerate(regions):
            put(s, r, 8 + k, f"=SUMIFS({DP('F')},{DP('A')},$A{r},{DP('C')},\"{g}\")", MWH)
            put(s, r, 8 + len(regions) + k, f"=SUMIFS({DP('G')},{DP('A')},$A{r},{DP('C')},\"{g}\")", MWH)
    last_c = 7 + 2 * len(regions)
    put(s, T_TOT, 1, "Total")
    for c in range(3, last_c + 1):
        put(s, T_TOT, c, f"=SUM({col(c)}{FIRST}:{col(c)}{TL})", INT if c == 3 else MWH)
    total_row(s, T_TOT, 1, last_c)
    s.freeze_panes = "B5"
    s.auto_filter.ref = f"A4:{col(last_c)}{TL}"
    # monthly table to the right
    m0 = last_c + 2
    put(s, 3, m0, "Per month", font=F_SECTION)
    header(s, 4, ["Month", "Days", "AS Up (MWh)", "AS Down (MWh)", "Net Up-Down (MWh)", "Gross (MWh)",
                  "Avg Up per day (MWh)", "Avg Down per day (MWh)"], [9, 7, 13, 13, 13, 13, 12, 12], c0=m0)
    dt_rng = lambda c: ref("Daily Totals", c, FIRST, TL)
    for i, m in enumerate(months):
        r = FIRST + i
        put(s, r, m0, m)
        put(s, r, m0 + 1, f"=COUNTIF({dt_rng('B')},{col(m0)}{r})", INT)
        put(s, r, m0 + 2, f"=SUMIFS({dt_rng('D')},{dt_rng('B')},{col(m0)}{r})", MWH)
        put(s, r, m0 + 3, f"=SUMIFS({dt_rng('E')},{dt_rng('B')},{col(m0)}{r})", MWH)
        put(s, r, m0 + 4, f"={col(m0 + 2)}{r}-{col(m0 + 3)}{r}", MWH)
        put(s, r, m0 + 5, f"={col(m0 + 2)}{r}+{col(m0 + 3)}{r}", MWH)
        put(s, r, m0 + 6, f"=IF({col(m0 + 1)}{r}=0,0,{col(m0 + 2)}{r}/{col(m0 + 1)}{r})", MWH)
        put(s, r, m0 + 7, f"=IF({col(m0 + 1)}{r}=0,0,{col(m0 + 3)}{r}/{col(m0 + 1)}{r})", MWH)
    mr = FIRST + len(months)
    put(s, mr, m0, "Total")
    for c in range(m0 + 1, m0 + 6):
        put(s, mr, c, f"=SUM({col(c)}{FIRST}:{col(c)}{mr - 1})", INT if c == m0 + 1 else MWH)
    put(s, mr, m0 + 6, f"=IF({col(m0 + 1)}{mr}=0,0,{col(m0 + 2)}{mr}/{col(m0 + 1)}{mr})", MWH)
    put(s, mr, m0 + 7, f"=IF({col(m0 + 1)}{mr}=0,0,{col(m0 + 3)}{mr}/{col(m0 + 1)}{mr})", MWH)
    total_row(s, mr, m0, m0 + 7)
    month_rows = (FIRST, mr - 1)

    # ---------------------------------------------------------------- Matrix Up / Matrix Down (data)
    piv = {k: day.pivot_table(index="plant", columns="date", values=v, aggfunc="sum")
           for k, v in (("Up", "up_mwh"), ("Down", "down_mwh"))}
    for kind, colour in (("Up", UP_HEX), ("Down", DOWN_HEX)):
        s = ws[f"Matrix {kind}"]
        titled(s, f"AS {kind} (MWh) - every plant x every day, in one place",
               f"Values from 'Daily by Plant' (blank = no AS {kind} that day). Rows in the same order as "
               f"'Plant Summary'. Darker = more MWh.")
        header(s, 4, ["Region", "Plant", "Technology", f"Total AS {kind} (MWh)"], [7, 26, 22, 13])
        for j, d in enumerate(dates):
            put(s, 4, 5 + j, d, "dd-mmm", font=F_HEAD, fill=FILL_HEAD, align=HEAD_ALIGN)
            s.column_dimensions[col(5 + j)].width = 7.5
        pv = piv[kind].reindex(index=plants["plant"], columns=[pd.Timestamp(d) for d in dates])
        last = col(4 + n_dt)
        for i, (p, vals) in enumerate(zip(plants.itertuples(index=False), pv.to_numpy())):
            r = FIRST + i
            s.cell(row=r, column=1, value=p.region)
            s.cell(row=r, column=2, value=p.plant)
            s.cell(row=r, column=3, value=p.technology)
            put(s, r, 4, f"=SUM(E{r}:{last}{r})", MWH0, font=F_BOLD)
            for j, v in enumerate(vals):
                if v == v and v > 0:                 # skip NaN and zero
                    put(s, r, 5 + j, round(float(v), 3), MWH0)
        put(s, P_TOT, 2, "All plants")
        for c in range(4, 5 + n_dt):
            put(s, P_TOT, c, f"=SUM({col(c)}{FIRST}:{col(c)}{PL})", MWH0)
        total_row(s, P_TOT, 1, 4 + n_dt)
        s.conditional_formatting.add(f"E{FIRST}:{last}{PL}",
                                     ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                    end_type="percentile", end_value=99, end_color=colour))
        s.freeze_panes = "E5"
        s.auto_filter.ref = f"A4:{last}{PL}"

    # ---------------------------------------------------------------- Time Slots
    s = ws["Time Slots"]
    titled(s, "In which time slots is AS scheduled? Whole date range, all plants",
           "Formulas over 'Hourly Data'. Avg MW = MWh in that hour / number of days. "
           "For your own date range, region or plant use 'Range Analyser'.")
    hour_labels = [f"{h:02d}:00-{h + 1:02d}:00" for h in range(24)]

    def hour_table(sheet, top: int, up_f, dn_f, days_cell: str) -> int:
        """24 hourly rows from top+1; returns the total row."""
        header(sheet, top, ["Hour slot", "AS Up (MWh)", "Share of Up", "AS Down (MWh)", "Share of Down",
                            "Gross (MWh)", "Share of Gross", "Avg Up (MW)", "Avg Down (MW)"],
               [16, 13, 10, 13, 10, 13, 10, 11, 11])
        tot = top + 25
        for h in range(24):
            r = top + 1 + h
            put(sheet, r, 1, hour_labels[h])
            put(sheet, r, 2, up_f(h), MWH)
            put(sheet, r, 3, f"=IF($B${tot}=0,0,B{r}/$B${tot})", PCT)
            put(sheet, r, 4, dn_f(h), MWH)
            put(sheet, r, 5, f"=IF($D${tot}=0,0,D{r}/$D${tot})", PCT)
            put(sheet, r, 6, f"=B{r}+D{r}", MWH)
            put(sheet, r, 7, f"=IF($F${tot}=0,0,F{r}/$F${tot})", PCT)
            put(sheet, r, 8, f"=IF({days_cell}=0,0,B{r}/{days_cell})", MWH)
            put(sheet, r, 9, f"=IF({days_cell}=0,0,D{r}/{days_cell})", MWH)
        put(sheet, tot, 1, "Total")
        for c, L in ((2, "B"), (4, "D"), (6, "F")):
            put(sheet, tot, c, f"=SUM({L}{top + 1}:{L}{top + 24})", MWH)
        for c, L in ((3, "C"), (5, "E"), (7, "G")):
            put(sheet, tot, c, f"=SUM({L}{top + 1}:{L}{top + 24})", PCT)
        put(sheet, tot, 8, f"=IF({days_cell}=0,0,B{tot}/{days_cell}/24)", MWH)
        put(sheet, tot, 9, f"=IF({days_cell}=0,0,D{tot}/{days_cell}/24)", MWH)
        total_row(sheet, tot, 1, 9)
        return tot

    def band_table(sheet, top: int, hour_top: int) -> int:
        """Band rows summing the hour rows that start at hour_top+1; returns the total row."""
        put(sheet, top - 1, 1, "Time bands", font=F_SECTION)
        header(sheet, top, ["Time band", "Hours", "AS Up (MWh)", "Share of Up", "AS Down (MWh)",
                            "Share of Down", "Gross (MWh)", "Share of Gross"])
        tot = top + len(bands) + 1
        for k, (name, hrs) in enumerate(bands):
            r = top + 1 + k
            put(sheet, r, 1, name)
            put(sheet, r, 2, len(hrs), INT)
            put(sheet, r, 3, "=" + "+".join(f"B{hour_top + 1 + h}" for h in hrs), MWH)
            put(sheet, r, 4, f"=IF($C${tot}=0,0,C{r}/$C${tot})", PCT)
            put(sheet, r, 5, "=" + "+".join(f"D{hour_top + 1 + h}" for h in hrs), MWH)
            put(sheet, r, 6, f"=IF($E${tot}=0,0,E{r}/$E${tot})", PCT)
            put(sheet, r, 7, f"=C{r}+E{r}", MWH)
            put(sheet, r, 8, f"=IF($G${tot}=0,0,G{r}/$G${tot})", PCT)
        put(sheet, tot, 1, "Total")
        for c in range(2, 9):
            L = col(c)
            put(sheet, tot, c, f"=SUM({L}{top + 1}:{L}{tot - 1})", PCT if c in (4, 6, 8) else
                (INT if c == 2 else MWH))
        total_row(sheet, tot, 1, 8)
        return tot

    put(s, 31, 1, "Days in data")
    put(s, 31, 2, f"=COUNT({ref('Daily Totals', 'A', FIRST, TL)})", INT, font=F_BOLD)
    hour_table(s, 4, lambda h: f"=SUM({HD(hour_col(h))})", lambda h: f"=SUM({HD(hour_col(h, False))})", "$B$31")
    band_table(s, 34, 4)                              # band names land in A35:A38 (used by Plant Summary)
    # month x hour heatmaps
    hm_top = 34 + len(bands) + 4
    for kind, up, colour in (("Up", True, UP_HEX), ("Down", False, DOWN_HEX)):
        put(s, hm_top - 1, 1, f"AS {kind} (MWh) by month and hour - darker = more", font=F_SECTION)
        header(s, hm_top, ["Month"] + [f"{h:02d}" for h in range(24)] + ["Total"])
        for i, m in enumerate(months):
            r = hm_top + 1 + i
            put(s, r, 1, m)
            for h in range(24):
                put(s, r, 2 + h, f"=SUMIFS({HD(hour_col(h, up))},{HD('B')},$A{r})", MWH0)
            put(s, r, 26, f"=SUM(B{r}:Y{r})", MWH0, font=F_BOLD)
        s.conditional_formatting.add(f"B{hm_top + 1}:Y{hm_top + len(months)}",
                                     ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                    end_type="max", end_color=colour))
        hm_top += len(months) + 4
    for c in range(10, 27):
        s.column_dimensions[col(c)].width = max(s.column_dimensions[col(c)].width or 0, 9)
    s.freeze_panes = "A5"
    cats = Reference(s, min_col=1, min_row=5, max_row=28)
    s.add_chart(bar_chart("AS energy by hour of day (whole range)", "MWh", cats,
                          [(Reference(s, min_col=2, min_row=5, max_row=28), "AS Up", UP_HEX),
                           (Reference(s, min_col=4, min_row=5, max_row=28), "AS Down", DOWN_HEX)]), "K4")
    bcats = Reference(s, min_col=1, min_row=35, max_row=34 + len(bands))
    s.add_chart(bar_chart("AS energy by time band (whole range)", "MWh", bcats,
                          [(Reference(s, min_col=3, min_row=35, max_row=34 + len(bands)), "AS Up", UP_HEX),
                           (Reference(s, min_col=5, min_row=35, max_row=34 + len(bands)), "AS Down", DOWN_HEX)],
                          width=16, height=7.5), "K22")

    # ---------------------------------------------------------------- Lists (hidden, for drop-downs)
    s = ws["Lists"]
    for i, v in enumerate(["All"] + regions):
        s.cell(row=1 + i, column=1, value=v)
    for i, v in enumerate(["All"] + techs):
        s.cell(row=1 + i, column=2, value=v)
    plant_list = ["All"] + sorted(plants["plant"], key=str.lower)
    for i, v in enumerate(plant_list):
        s.cell(row=1 + i, column=3, value=v)
    s.sheet_state = "hidden"

    # ---------------------------------------------------------------- Range Analyser
    s = ws["Range Analyser"]
    titled(s, "Range Analyser - pick a date range (and optionally a region, technology or plant)",
           "Change only the yellow cells. Everything below recalculates from 'Hourly Data' and 'Daily by Plant'.")
    inputs = [("From date", dates[0], DATE), ("To date", dates[-1], DATE), ("Region", "All", None),
              ("Technology", "All", None), ("Plant", "All", None)]
    for i, (label, v, fmt) in enumerate(inputs):
        r = 4 + i
        put(s, r, 1, label, font=F_BOLD)
        put(s, r, 2, v, fmt, font=F_INPUT, fill=FILL_INPUT)
    put(s, 4, 3, f"Data available: {dates[0]:%d-%b-%Y} to {dates[-1]:%d-%b-%Y}", font=F_NOTE)
    put(s, 6, 3, "All, or one region", font=F_NOTE)
    put(s, 7, 3, "All, or one technology", font=F_NOTE)
    put(s, 8, 3, "All, or one plant (pick from the list)", font=F_NOTE)
    put(s, 3, 4, "criteria used", font=F_NOTE)
    for r in (6, 7, 8):
        put(s, r, 4, f'=IF(OR(B{r}="All",B{r}=""),"*",B{r})', font=F_NOTE)
    dv_date = DataValidation(type="date", operator="between", formula1=f"DATE({dates[0]:%Y,%m,%d})".replace(",0", ","),
                             formula2=f"DATE({dates[-1]:%Y,%m,%d})".replace(",0", ","), allow_blank=False,
                             showErrorMessage=True, errorTitle="Date outside the data",
                             error=f"Pick a date between {dates[0]:%d-%b-%Y} and {dates[-1]:%d-%b-%Y}.")
    dv_reg = DataValidation(type="list", formula1=f"Lists!$A$1:$A${1 + len(regions)}", showErrorMessage=True)
    dv_typ = DataValidation(type="list", formula1=f"Lists!$B$1:$B${1 + len(techs)}", showErrorMessage=True)
    dv_pl = DataValidation(type="list", formula1=f"Lists!$C$1:$C${len(plant_list)}", showErrorMessage=True)
    for dv, cells in ((dv_date, "B4:B5"), (dv_reg, "B6"), (dv_typ, "B7"), (dv_pl, "B8")):
        s.add_data_validation(dv)
        dv.add(cells)
    crit = (f"{HD('A')},\">=\"&$B$4,{HD('A')},\"<=\"&$B$5,{HD('C')},$D$6,{HD('E')},$D$7,{HD('D')},$D$8")
    dcrit = (f"{DP('A')},\">=\"&$B$4,{DP('A')},\"<=\"&$B$5,{DP('C')},$D$6,{DP('E')},$D$7,{DP('D')},$D$8")
    kpis = [("Days with data in range", f"=COUNTIFS({ref('Daily Totals', 'A', FIRST, TL)},\">=\"&B4,"
                                         f"{ref('Daily Totals', 'A', FIRST, TL)},\"<=\"&B5)", INT),
            ("Plant-days with AS", f"=COUNTIFS({dcrit})", INT),
            ("AS Up (MWh)", "=B43", MWH), ("AS Down (MWh)", "=D43", MWH),
            ("Net Up-Down (MWh)", "=B43-D43", MWH), ("Gross Up+Down (MWh)", "=F43", MWH)]
    for i, (label, f, fmt) in enumerate(kpis):
        put(s, 10 + i, 1, label, font=F_BOLD)
        put(s, 10 + i, 2, f, fmt, font=F_BOLD)
    put(s, 17, 1, "By hour", font=F_SECTION)
    hour_table(s, 18, lambda h: f"=SUMIFS({HD(hour_col(h))},{crit})",
               lambda h: f"=SUMIFS({HD(hour_col(h, False))},{crit})", "$B$10")
    band_table(s, 46, 18)
    cats = Reference(s, min_col=1, min_row=19, max_row=42)
    s.add_chart(bar_chart("AS energy by hour - selected range", "MWh", cats,
                          [(Reference(s, min_col=2, min_row=19, max_row=42), "AS Up", UP_HEX),
                           (Reference(s, min_col=4, min_row=19, max_row=42), "AS Down", DOWN_HEX)]), "K4")
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 16
    s.column_dimensions["D"].width = 14

    # ---------------------------------------------------------------- Data Check
    s = ws["Data Check"]
    files, check = t["files"], t["check"]
    titled(s, "Data check - which files were used, and does every region balance?",
           "Checksum: in every 15-min block, the plants' AS plus the region's VAE row must add to zero. "
           "Status OK = within 0.5 MW in all 96 blocks.")
    n_bad = int((check["status"] != "OK").sum())
    cover = [("First date", dates[0], DATE), ("Last date", dates[-1], DATE),
             ("Calendar days in range", (dates[-1] - dates[0]).days + 1, INT), ("Files used", len(files), INT),
             ("Problems found while reading", "; ".join(t.get("problems") or []) or "None", None),
             ("VAE checksum", f"{len(check) - n_bad} of {len(check)} date-regions OK", None),
             ("Schedule types add up to Total",
              f"{int((sched['check'].abs() <= 0.5).sum())} of {len(sched)} plant-days (within 0.5 MWh)", None),
             ("Late revisions", "WBES keeps revising a date for several days; the last ~7 dates may still change. "
                                "Re-run wbes_pull.py, then this report.", None)]
    for i, (label, v, fmt) in enumerate(cover):
        put(s, 4 + i, 1, label, font=F_BOLD)
        put(s, 4 + i, 2, v, fmt)
    top = 13
    put(s, top - 1, 1, "Files used (latest revision of each date)", font=F_SECTION)
    header(s, top, ["Date", "File", "Revision", "Exported at", "Region sheets", "Plants with AS",
                    "AS Up (MWh)", "AS Down (MWh)"], [12, 34, 9, 20, 16, 9, 13, 13])
    for i, f in enumerate(files.sort_values("date").itertuples(index=False)):
        r = top + 1 + i
        put(s, r, 1, f.date.to_pydatetime(), DATE)
        s.cell(row=r, column=2, value=f.file)
        put(s, r, 3, f.revision, INT)
        s.cell(row=r, column=4, value=f.exported_at)
        s.cell(row=r, column=5, value=f.regions)
        put(s, r, 6, int(f.plants) if f.plants == f.plants else 0, INT)
        put(s, r, 7, float(f.up) if f.up == f.up else 0.0, MWH)
        put(s, r, 8, float(f.down) if f.down == f.down else 0.0, MWH)
    c0 = 10
    put(s, top - 1, c0, "VAE checksum per date and region", font=F_SECTION)
    header(s, top, ["Date", "Region", "Plants with AS", "VAE row present", "Max |plants + VAE| (MW)", "Status"],
           [12, 8, 9, 9, 13, 8], c0=c0)
    for i, k in enumerate(check.itertuples(index=False)):
        r = top + 1 + i
        put(s, r, c0, k.date.to_pydatetime(), DATE)
        s.cell(row=r, column=c0 + 1, value=k.region)
        put(s, r, c0 + 2, int(k.plants), INT)
        s.cell(row=r, column=c0 + 3, value="Yes" if k.vae_row else "No")
        put(s, r, c0 + 4, float(k.max_diff_mw), "0.000")
        s.cell(row=r, column=c0 + 5, value=k.status)
    s.freeze_panes = f"A{top + 1}"

    # ---------------------------------------------------------------- Plants (technology + capacity)
    s = ws["Plants"]
    titled(s, "Plant list - technology, how it was identified, and capacity",
           "Technology: automatic from the plant's own WBES 'Total' schedule shape and its code, unless set in "
           "plant_master.csv. Max scheduled MW = highest MW WBES scheduled in the period (not nameplate). "
           "Nameplate MW = web-checked, high/medium confidence only; see Source and Note.")
    heads = ["Region", "Plant", "Technology", "How identified", "Share of output at night (19:00-05:30)",
             "Charging / injection ratio", "WBES plant type", "Max scheduled MW (WBES)", "Max charging MW (WBES)",
             "Nameplate MW (web)", "AC / DC", "BESS MWh (web)", "Project (web)", "Developer (web)", "Source (web)",
             "Confidence (web)", "Note", "AS Up (MWh)", "AS Down (MWh)"]
    header(s, 4, heads, [8, 26, 22, 48, 11, 10, 20, 11, 11, 11, 7, 9, 30, 22, 40, 10, 60, 12, 12])
    txt = lambda v: v if isinstance(v, str) and v else None
    for i, p in enumerate(plants.itertuples(index=False)):
        r = FIRST + i
        for c, v in ((1, p.region), (2, p.plant), (3, p.technology), (4, p.tech_evidence), (7, p.type),
                     (11, txt(p.ac_or_dc)), (13, txt(p.project_name)), (14, txt(p.developer)), (15, txt(p.source_url)),
                     (16, txt(p.confidence)), (17, txt(p.note))):
            s.cell(row=r, column=c, value=v)
        put(s, r, 5, _num(p.night_share), PCT)
        put(s, r, 6, _num(p.charge_ratio), PCT)
        put(s, r, 8, _num(p.max_sched_mw), MWH)
        put(s, r, 9, _num(p.max_charge_mw), MWH)
        put(s, r, 10, _num(p.nameplate_mw), MWH)
        put(s, r, 12, _num(p.bess_mwh), MWH)
        put(s, r, 18, f"=SUMIFS({DP('F')},{DP('D')},$B{r})", MWH)
        put(s, r, 19, f"=SUMIFS({DP('G')},{DP('D')},$B{r})", MWH)
    s.freeze_panes = "C5"
    s.auto_filter.ref = f"A4:S{PL}"
    PLR = lambda c: ref("Plants", c, FIRST, PL)
    c0 = 21
    put(s, 3, c0, "By technology", font=F_SECTION)
    header(s, 4, ["Technology", "Plants", "Max scheduled MW (sum)", "Plants with nameplate", "Nameplate MW (sum)",
                  "AS Up (MWh)", "AS Down (MWh)"], [24, 8, 12, 10, 12, 13, 13], c0=c0)
    for k, x in enumerate(techs):
        r = FIRST + k
        put(s, r, c0, x)
        put(s, r, c0 + 1, f"=COUNTIF({PLR('C')},{col(c0)}{r})", INT)
        put(s, r, c0 + 2, f"=SUMIFS({PLR('H')},{PLR('C')},{col(c0)}{r})", MWH0)
        put(s, r, c0 + 3, f"=COUNTIFS({PLR('C')},{col(c0)}{r},{PLR('J')},\">0\")", INT)
        put(s, r, c0 + 4, f"=SUMIFS({PLR('J')},{PLR('C')},{col(c0)}{r})", MWH0)
        put(s, r, c0 + 5, f"=SUMIFS({PLR('R')},{PLR('C')},{col(c0)}{r})", MWH0)
        put(s, r, c0 + 6, f"=SUMIFS({PLR('S')},{PLR('C')},{col(c0)}{r})", MWH0)
    r = FIRST + len(techs)
    put(s, r, c0, "Total")
    for k in range(1, 7):
        L = col(c0 + k)
        put(s, r, c0 + k, f"=SUM({L}{FIRST}:{L}{r - 1})", INT if k in (1, 3) else MWH0)
    total_row(s, r, c0, c0 + 6)

    # ---------------------------------------------------------------- Summary
    s = ws["Summary"]
    titled(s, "AS (Ancillary Services) schedule - summary",
           f"Source: wbes_ns_data Net Schedule files, {dates[0]:%d-%b-%Y} to {dates[-1]:%d-%b-%Y} "
           f"({n_dt} days, latest revision of each date). Generated {t['generated']:%d-%b-%Y %H:%M}. "
           f"All numbers below are formulas.")
    for c, w in zip("ABCDEF", [40, 26, 16, 16, 14, 3]):
        s.column_dimensions[c].width = w
    ts = lambda c, r1, r2: ref("Time Slots", c, r1, r2)
    up_tot, dn_tot = f"'Plant Summary'!$H${P_TOT}", f"'Plant Summary'!$I${P_TOT}"
    put(s, 4, 1, "Headline numbers", font=F_SECTION)
    kp = [("Date range", f"=TEXT(MIN({ref('Daily Totals', 'A', FIRST, TL)}),\"dd-mmm-yyyy\")&\" to \"&"
                         f"TEXT(MAX({ref('Daily Totals', 'A', FIRST, TL)}),\"dd-mmm-yyyy\")", None, None),
          ("Days of data", f"=COUNT({ref('Daily Totals', 'A', FIRST, TL)})", INT, None),
          ("Plants with any AS", f"=COUNTIF({PS('E')},\">0\")", INT, None),
          ("Plant-days with AS", f"=COUNT({DP('A')})", INT, None),
          ("Total AS Up (MWh)", f"={up_tot}", MWH0, ("=B9/1000", MU)),
          ("Total AS Down (MWh)", f"={dn_tot}", MWH0, ("=B10/1000", MU)),
          ("Gross Up+Down (MWh)", "=B9+B10", MWH0, ("=B11/1000", MU)),
          ("Net Up-Down (MWh)  (negative = more Down than Up)", "=B9-B10", MWH0, ("=B12/1000", MU)),
          ("Average AS Up per day (MWh/day)", "=IF(B6=0,0,B9/B6)", MWH0, None),
          ("Average AS Down per day (MWh/day)", "=IF(B6=0,0,B10/B6)", MWH0, None),
          ("Hour with the most AS Up  (share of all Up)",
           f"=INDEX({ts('A', 5, 28)},MATCH(MAX({ts('B', 5, 28)}),{ts('B', 5, 28)},0))", None,
           (f"=MAX({ts('C', 5, 28)})", PCT)),
          ("Hour with the most AS Down  (share of all Down)",
           f"=INDEX({ts('A', 5, 28)},MATCH(MAX({ts('D', 5, 28)}),{ts('D', 5, 28)},0))", None,
           (f"=MAX({ts('E', 5, 28)})", PCT)),
          ("Time band with the most AS Up  (share of all Up)",
           f"=INDEX({ts('A', 35, 34 + nb)},MATCH(MAX({ts('C', 35, 34 + nb)}),{ts('C', 35, 34 + nb)},0))", None,
           (f"=MAX({ts('D', 35, 34 + nb)})", PCT)),
          ("Time band with the most AS Down  (share of all Down)",
           f"=INDEX({ts('A', 35, 34 + nb)},MATCH(MAX({ts('E', 35, 34 + nb)}),{ts('E', 35, 34 + nb)},0))", None,
           (f"=MAX({ts('F', 35, 34 + nb)})", PCT))]
    for i, (label, f, fmt, extra) in enumerate(kp):
        r = 5 + i
        put(s, r, 1, label)
        put(s, r, 2, f, fmt, font=F_BOLD)
        if extra:
            put(s, r, 3, extra[0], extra[1])
    # top-10 lists
    r0 = 21
    for kind, key, val, tot in (("Up", "T", "H", up_tot), ("Down", "U", "I", dn_tot)):
        put(s, r0, 1, f"Top 10 plants by AS {kind} (whole range)", font=F_SECTION)
        header(s, r0 + 1, ["Plant", "Region / technology", "Rank", f"AS {kind} (MWh)", f"Share of all {kind}"])
        for k in range(1, 11):
            r = r0 + 1 + k
            m = f"MATCH(LARGE({PS(key)},{k}),{PS(key)},0)"
            put(s, r, 1, f"=INDEX({PS('C')},{m})")
            put(s, r, 2, f"=INDEX({PS('B')},{m})&\" / \"&INDEX({PS('D')},{m})")
            put(s, r, 3, k, INT)
            put(s, r, 4, f"=INDEX({PS(val)},{m})", MWH0)
            put(s, r, 5, f"=IF({tot}=0,0,D{r}/{tot})", PCT)
        r0 += 14
    # by region / type / month
    def split_table(top: int, title: str, label: str, items: list[str], crit_col: str) -> int:
        put(s, top, 1, title, font=F_SECTION)
        header(s, top + 1, [label, "AS Up (MWh)", "Share of Up", "AS Down (MWh)", "Share of Down"])
        tot = top + 2 + len(items)
        for k, it in enumerate(items):
            r = top + 2 + k
            put(s, r, 1, it)
            put(s, r, 2, f"=SUMIFS({DP('F')},{DP(crit_col)},$A{r})", MWH0)
            put(s, r, 3, f"=IF($B${tot}=0,0,B{r}/$B${tot})", PCT)
            put(s, r, 4, f"=SUMIFS({DP('G')},{DP(crit_col)},$A{r})", MWH0)
            put(s, r, 5, f"=IF($D${tot}=0,0,D{r}/$D${tot})", PCT)
        put(s, tot, 1, "Total")
        for c, L, fmt in ((2, "B", MWH0), (3, "C", PCT), (4, "D", MWH0), (5, "E", PCT)):
            put(s, tot, c, f"=SUM({L}{top + 2}:{L}{tot - 1})", fmt)
        total_row(s, tot, 1, 5)
        return tot + 2
    nxt = split_table(r0, "By region", "Region", regions, "C")
    nxt = split_table(nxt, "By technology", "Technology", techs, "E")
    nxt = split_table(nxt, "By WBES plant type", "WBES plant type", types, type_col)
    month_top = nxt
    nxt = split_table(nxt, "By month", "Month", months, "B")
    put(s, nxt, 1, "Whole schedule of the AS plants, by schedule type (all days, + = into grid)", font=F_SECTION)
    header(s, nxt + 1, ["Schedule type", "MWh", "MU", "Share of Total"])
    mtot = nxt + 1 + stypes.index("Total") + 1 if "Total" in stypes else None
    for k, x in enumerate(stypes):
        r = nxt + 2 + k
        put(s, r, 1, x, font=F_BOLD if x == "Total" else None)
        put(s, r, 2, f"=SUM({SB(st_col[x])})", MWH0)
        put(s, r, 3, f"=B{r}/1000", MU)
        if mtot:
            put(s, r, 4, f"=IF(ABS($B${mtot})<0.000001,0,B{r}/$B${mtot})", PCT)
    if mtot:
        total_row(s, mtot, 1, 4)
    nxt += len(stypes) + 3
    # key findings (static text, written from the numbers at generation time)
    put(s, nxt, 1, f"Key findings (text written from the data on {t['generated']:%d-%b-%Y}; "
                   f"the tables above update by formula, this text does not)", font=F_SECTION)
    for i, line in enumerate(findings(t, bands)):
        c = put(s, nxt + 1 + i, 1, f"{i + 1}. {line}", align=WRAP_TOP)
        s.merge_cells(start_row=nxt + 1 + i, start_column=1, end_row=nxt + 1 + i, end_column=5)
        s.row_dimensions[nxt + 1 + i].height = 42
    # charts
    s.add_chart(bar_chart("AS energy by hour of day", "MWh", Reference(ws["Time Slots"], min_col=1, min_row=5, max_row=28),
                          [(Reference(ws["Time Slots"], min_col=2, min_row=5, max_row=28), "AS Up", UP_HEX),
                           (Reference(ws["Time Slots"], min_col=4, min_row=5, max_row=28), "AS Down", DOWN_HEX)]),
                "H4")
    s.add_chart(line_chart("AS energy per day, all plants", "MWh per day",
                           Reference(ws["Daily Totals"], min_col=1, min_row=FIRST, max_row=TL),
                           [(Reference(ws["Daily Totals"], min_col=4, min_row=FIRST, max_row=TL), "AS Up", UP_HEX),
                            (Reference(ws["Daily Totals"], min_col=5, min_row=FIRST, max_row=TL), "AS Down",
                             DOWN_HEX)]), "H23")
    s.add_chart(bar_chart("AS energy per month", "MWh",
                          Reference(s, min_col=1, min_row=month_top + 2, max_row=month_top + 1 + len(months)),
                          [(Reference(s, min_col=2, min_row=month_top + 2, max_row=month_top + 1 + len(months)),
                            "AS Up", UP_HEX),
                           (Reference(s, min_col=4, min_row=month_top + 2, max_row=month_top + 1 + len(months)),
                            "AS Down", DOWN_HEX)], width=16), "H42")

    # ---------------------------------------------------------------- Read Me
    s = ws["Read Me"]
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 110
    put(s, 1, 1, "AS (Ancillary Services) schedule analysis - read me", font=F_TITLE)
    lines = [
        ("What this answers", None),
        ("1. Day-wise per plant", "How much AS each plant was scheduled each day, Up and Down, and in which time "
                                  "slots -> 'Daily by Plant' (one row per plant per day) and 'Matrix Up' / "
                                  "'Matrix Down' (every plant x every day on one sheet)."),
        ("2. Cumulative per plant", "Total AS of each plant over the whole range -> 'Plant Summary'."),
        ("3. Time slots", "Share and MWh of AS in each hour and time band -> 'Time Slots' (whole range), "
                          "'Plant Time Bands' (per plant), 'Range Analyser' (your own dates / region / plant)."),
        ("4. Brief analysis", "'Summary': headline numbers, top plants, region / technology / month split, "
                              "schedule mix, key findings."),
        ("5. Technology + capacity", "'Plants': solar / wind / hybrid part / BESS / conventional for every plant, "
                                     "how it was identified, max scheduled MW, web-checked nameplate MW + source."),
        ("6. Whole schedule", "'Schedule by Type': every day of every AS plant, MWh per schedule type (ISGS, URS, "
                              "OA_GNA, OA_TGNA, OA_PX, OA_REMC, AS, SCED, SCUC, Total). Totals per plant on "
                              "'Plant Summary' (right-hand columns)."),
        ("", None),
        ("Source and coverage", None),
        ("Files", f"wbes_pipeline/wbes_ns_data/<yyyy-mm>/NetSchedule_<date>_revNNN.xlsx, {len(files)} files, "
                  f"{dates[0]:%d-%b-%Y} to {dates[-1]:%d-%b-%Y}. Each file is the latest WBES revision of that "
                  f"date at download time. Only rows with schedule type 'AS' are used. See 'Data Check'."),
        ("Refresh", "In the wbes_pipeline folder run:  .venv\\Scripts\\python.exe as_report.py   "
                    "(optional: --from 2026-06-01 --to 2026-09-30). Takes about 2 minutes."),
        ("", None),
        ("How the numbers are read", None),
        ("Units", "WBES gives MW per 15-min block (average MW in that block). Energy in MWh = MW x 0.25 h. "
                  "1 MU = 1,000 MWh. Confirmed: a 2,400 MW plant (KUDGI) shows ~1,492 MW in its ISGS row, "
                  "which only makes sense as MW."),
        ("Sign: Up vs Down", "WBES sign convention: drawal positive, injection negative (a generator's ISGS row is "
                             "negative). So an AS value below zero = extra injection = AS UP; above zero = backing "
                             "down = AS DOWN. Both are shown here as positive MWh. Supported by the data: AS Down is "
                             "concentrated 09:00-14:00 (solar surplus hours), as expected."),
        ("Net and Gross", "Net = Up - Down (negative = more backing down than extra injection). Gross = Up + Down "
                          "(total AS activity). A plant can have both Up and Down on the same day."),
        ("VAE rows", "VAE_ER, VAE_WR, VAE_NR, VAE_SR, VAE_NER are the Virtual Ancillary Entities: the regional "
                     "counter-party of every AS schedule. In every block VAE = minus the sum of all plants in that "
                     "region. They are excluded from plant figures and used as a checksum ('Data Check')."),
        ("WBES plant type", "From WBES utility_type_id: 4 = ISGS / central generator, 16 = IPP / RE / storage "
                            "(other ISTS generators), 5 = State control area. Labels inferred (not documented by WBES)."),
        ("Technology", "From each plant's own WBES 'Total' schedule: Solar = under 2% of its output at night "
                       "(19:00-05:30); Wind = over 35% at night; BESS = code has BESS and it charges (draws) as well "
                       "as discharges; Hybrid part = code ends _HS/_HW, or the same code exists with S and W legs "
                       "(e.g. ASEJOPL_S_FTG2 + ASEJOPL_W_FTG2); Conventional = no renewable (OA_REMC) schedule, or a "
                       "known thermal plant. Correct any plant in plant_master.csv (column 'technology') and re-run."),
        ("Capacity", "Max scheduled MW = highest MW WBES scheduled the plant in the period: a floor, not nameplate "
                     "(solar is often scheduled at its contracted AC capacity). Nameplate MW = web-checked for the "
                     "~60 largest RE / BESS plants by AS, kept only where the match is high or medium confidence; "
                     "low-confidence leads are in the Note. Source URL given for every number."),
        ("Schedule types", "MWh per type per day = sum of the 96 block MW x 0.25, sign flipped so + = into the grid. "
                           "AS in 'Schedule by Type' is net (Up - Down); AS Up and Down separately are in "
                           "'Daily by Plant'."),
        ("Time bands", "; ".join(f"{n} (hours {', '.join(f'{h:02d}' for h in hrs)})" for n, hrs in bands)
                       + ". Assigned by the hour the 15-min block starts in. Change BANDS in as_report.py to re-cut."),
        ("Time slots text", "e.g. '18:00-22:30, 23:15-24:00' = the 15-min blocks in which that plant had AS Up (or "
                            "Down) that day, merged into continuous ranges."),
        ("", None),
        ("Limits", None),
        ("SRAS vs TRAS", "The WBES 'AS' row combines all ancillary services; this data cannot split SRAS from TRAS."),
        ("Late revisions", "WBES revises a date for several days. The last ~7 dates may still change; re-run "
                           "wbes_pull.py and then this report."),
        ("Formulas", "Summary sheets are formulas; Excel computes them when the file opens (allow a few seconds). "
                     "Data sheets hold values read from the source files."),
    ]
    for i, (a, b) in enumerate(lines):
        r = 3 + i
        if b is None:
            put(s, r, 1, a, font=F_SECTION)
        else:
            put(s, r, 1, a, font=F_BOLD, align=WRAP_TOP)
            put(s, r, 2, b, align=WRAP_TOP)
    put(s, 4 + len(lines), 1, "Sheets", font=F_SECTION)
    for i, name in enumerate(sheets[:-1]):
        put(s, 5 + len(lines) + i, 1, name, font=F_BOLD)

    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    try:
        tmp.replace(path)
    except PermissionError:                           # the old report is open in Excel: keep it, save beside it
        path = path.with_name(f"{path.stem} (updated {dt.datetime.now():%H%M}).xlsx")
        tmp.replace(path)
        print(f"Note: the previous report is open in Excel, so this one was saved as {path.name}")
    return path


# ----------------------------------------------------------------------------- findings text
def findings(t: dict, bands: list) -> list[str]:
    day, plants = t["day"], t["plants"]
    n_days = len(t["dates"])
    up, dn = day["up_mwh"].sum(), day["down_mwh"].sum()
    hu = day[[f"up_h{h:02d}" for h in range(24)]].sum().to_numpy()
    hd = day[[f"down_h{h:02d}" for h in range(24)]].sum().to_numpy()
    band_up = {n: hu[hrs].sum() / up for n, hrs in bands} if up else {}
    band_dn = {n: hd[hrs].sum() / dn for n, hrs in bands} if dn else {}
    bu, bd = max(band_up, key=band_up.get), max(band_dn, key=band_dn.get)
    pu = plants.sort_values("up", ascending=False).iloc[0]
    pdn = plants.sort_values("down", ascending=False).iloc[0]
    top10_dn = plants.sort_values("down", ascending=False).head(10)["down"].sum() / dn
    by_type = day.groupby("technology")[["up_mwh", "down_mwh"]].sum()
    by_reg = day.groupby("region")[["up_mwh", "down_mwh"]].sum()
    by_month = day.groupby("month")[["up_mwh", "down_mwh"]].sum()
    hh = lambda h: f"{h:02d}:00-{h + 1:02d}:00"
    t_dn = by_type["down_mwh"].idxmax()
    t_up = by_type["up_mwh"].idxmax()
    return [
        f"Over {n_days} days, {plants.shape[0]} plants were scheduled AS. Total AS Up = {up:,.0f} MWh "
        f"({up / 1000:,.1f} MU); total AS Down = {dn:,.0f} MWh ({dn / 1000:,.1f} MU). Down is "
        f"{dn / up:.1f} times Up, i.e. AS was used mostly to back plants down.",
        f"AS Down is concentrated in the {bd} band ({band_dn[bd]:.0%} of all Down). The single biggest hour is "
        f"{hh(int(hd.argmax()))} with {hd.max() / dn:.1%} of all Down.",
        f"AS Up is largest in the {bu} band ({band_up[bu]:.0%} of all Up); the biggest hour is "
        f"{hh(int(hu.argmax()))} with {hu.max() / up:.1%} of all Up.",
        f"Biggest AS Up plant: {pu.plant} ({pu.region}) with {pu.up:,.0f} MWh ({pu.up / up:.1%} of all Up). "
        f"Biggest AS Down plant: {pdn.plant} ({pdn.region}) with {pdn.down:,.0f} MWh ({pdn.down / dn:.1%} of all "
        f"Down). The top 10 Down plants carry {top10_dn:.0%} of all Down.",
        f"By technology: '{t_dn}' carries {by_type.loc[t_dn, 'down_mwh'] / dn:.0%} of AS Down; "
        f"'{t_up}' carries {by_type.loc[t_up, 'up_mwh'] / up:.0%} of AS Up.",
        f"By region: {by_reg['up_mwh'].idxmax()} has the most AS Up ({by_reg['up_mwh'].max() / up:.0%}); "
        f"{by_reg['down_mwh'].idxmax()} has the most AS Down ({by_reg['down_mwh'].max() / dn:.0%}).",
        f"By month: AS Down was highest in {by_month['down_mwh'].idxmax()} "
        f"({by_month['down_mwh'].max():,.0f} MWh); AS Up was highest in {by_month['up_mwh'].idxmax()} "
        f"({by_month['up_mwh'].max():,.0f} MWh). Note the first and last months are part-months.",
    ]
